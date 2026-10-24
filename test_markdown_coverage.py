import base64
import io
import unittest
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from app.converter import convert_md_to_docx

SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" width="320" height="100" viewBox="0 0 320 100">'
    '<rect width="320" height="100" fill="#6B3FC9"/></svg>'
)
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def data_uri(mime: str, data: bytes) -> str:
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"


def webp_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (15, 76, 92)).save(buf, format="WEBP")
    return buf.getvalue()


def convert(md: str) -> Document:
    return Document(BytesIO(convert_md_to_docx(md)))


def body_text(doc: Document) -> list[str]:
    """Top-level paragraphs only."""
    return [p.text for p in doc.paragraphs]


def all_text(doc: Document) -> list[str]:
    """Every paragraph in document order, including those inside boxes and
    tables, with text from hyperlinks too."""
    return ["".join(t.text or "" for t in p.iter(qn("w:t")))
            for p in doc.element.body.iter(qn("w:p"))]


def boxes(doc: Document):
    """Top-level one-cell tables: quotes, alerts and code blocks."""
    return [t for t in doc.tables if len(t.rows) == 1 and len(t.columns) == 1]


def box_text(table) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in table._tbl.iter(qn("w:p"))]


class MarkdownCoverageTests(unittest.TestCase):
    def test_svg_and_webp_images_are_embedded(self):
        doc = convert(
            f"# T\n\n![svg]({data_uri('image/svg+xml', SVG.encode())})\n\n"
            f"![webp]({data_uri('image/webp', webp_bytes())})\n"
        )
        self.assertEqual(len(doc.inline_shapes), 2)
        # Sized as a share of the 720px web column: 320px of the 16cm text width.
        self.assertAlmostEqual(doc.inline_shapes[0].width.cm, 16 * 320 / 720, places=1)

    def test_html_width_sizes_the_picture(self):
        doc = convert(f'# T\n\n<img src="{data_uri("image/svg+xml", SVG.encode())}" width="160">\n')
        self.assertAlmostEqual(doc.inline_shapes[0].width.cm, 16 * 160 / 720, places=1)

    def test_svg_cannot_read_local_files(self):
        svg = (
            '<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
            'width="10" height="10"><image xlink:href="file:///etc/passwd" width="10" height="10"/></svg>'
        )
        doc = convert(f"![x]({data_uri('image/svg+xml', svg.encode())})\n")
        # Renders (the reference is ignored) without touching the file.
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_image_caption_from_italic_line(self):
        doc = convert(f"# T\n\n![svg]({data_uri('image/svg+xml', SVG.encode())})\n\n*The caption.*\n")
        # Picture and caption share a box that can't split across pages.
        (figure,) = boxes(doc)
        picture, caption = figure.cell(0, 0).paragraphs
        self.assertEqual(len(picture._p.xpath(".//w:drawing")), 1)
        self.assertEqual(caption.text, "The caption.")
        self.assertEqual(caption.alignment, 1)  # centered
        self.assertIsNotNone(figure.rows[0]._tr.trPr.find(qn("w:cantSplit")))

    def test_strikethrough_and_escapes(self):
        doc = convert("# T\n\n~~gone~~ and \\*literal\\* stars\n")
        p = doc.paragraphs[1]
        self.assertEqual(p.text, "gone and *literal* stars")
        self.assertTrue(p.runs[0].font.strike)
        self.assertFalse(any(run.italic for run in p.runs))

    def test_highlight_sub_sup_and_kbd(self):
        doc = convert("# T\n\n==key== H~2~O x^2^ <kbd>Cmd</kbd>\n")
        runs = doc.paragraphs[1].runs
        by_text = {run.text: run for run in runs}
        self.assertIsNotNone(by_text["key"]._r.rPr.find(qn("w:shd")))
        self.assertEqual([r.text for r in runs if r.font.subscript], ["2"])
        self.assertEqual([r.text for r in runs if r.font.superscript], ["2"])
        self.assertIsNotNone(by_text["Cmd"]._r.rPr.find(qn("w:bdr")))

    def test_smart_quotes_and_dashes(self):
        doc = convert('# T\n\n"Quoted" -- and --- and `"code"`\n')
        self.assertEqual(doc.paragraphs[1].text, "“Quoted” – and — and \"code\"")

    def test_blockquote_is_a_box_with_a_left_bar(self):
        doc = convert("# T\n\n> Quoted\n>\n> # Heading inside\n\nAfter\n")
        (quote,) = boxes(doc)
        cell_borders = quote._tbl.tblPr.find(qn("w:tblBorders"))
        self.assertEqual(cell_borders.find(qn("w:left")).get(qn("w:val")), "single")
        self.assertEqual(box_text(quote), ["Quoted", "Heading inside"])
        self.assertIn("After", body_text(doc))

    def test_alert_is_a_titled_box(self):
        doc = convert("# T\n\n> [!WARNING]\n> Back up first.\n")
        (alert,) = boxes(doc)
        self.assertEqual(box_text(alert), ["Warning", "Back up first."])
        fill = alert.cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"))
        self.assertEqual(fill, "FFF8E5")

    def test_horizontal_rule_is_drawn(self):
        doc = convert("# T\n\nAbove\n\n---\n\nBelow\n")
        rule = doc.paragraphs[2]
        self.assertEqual(rule.text, "")
        self.assertIsNotNone(rule._p.pPr.find(qn("w:pBdr")))

    def test_task_list_uses_check_boxes(self):
        doc = convert("# T\n\n- [x] Done\n- [ ] Open\n")
        self.assertEqual(body_text(doc)[1:3], ["☑\tDone", "☐\tOpen"])

    def test_adjacent_code_blocks_stay_separate(self):
        doc = convert("# T\n\n```\none\n```\n\n```\ntwo\n```\n")
        self.assertEqual([box_text(b) for b in boxes(doc)], [["one"], ["two"]])

    def test_code_is_highlighted_and_labelled(self):
        doc = convert("# T\n\n```python\ndef f():\n    return 1\n```\n")
        (code,) = boxes(doc)
        self.assertEqual(box_text(code), ["python", "def f():", "    return 1"])
        colors = {r.font.color.rgb for r in code.cell(0, 0).paragraphs[1].runs if r.font.color.rgb}
        self.assertGreater(len(colors), 1)

    def test_footnotes_link_to_their_notes(self):
        doc = convert("# T\n\nNote.[^1] H<sub>2</sub>O\n\n[^1]: The note.\n")
        p = doc.paragraphs[1]
        link = p._p.find(qn("w:hyperlink"))
        self.assertEqual(link.get(qn("w:anchor")), "__fn_1")
        self.assertEqual("".join(t.text for t in link.iter(qn("w:t"))), "1")
        sub = [run for run in p.runs if run.font.subscript]
        self.assertEqual([run.text for run in sub], ["2"])
        self.assertIn("1\tThe note.", body_text(doc))
        names = [b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))]
        self.assertIn("__fn_1", names)

    def test_links_to_headings_jump_to_them(self):
        doc = convert("# T\n\nSee [the tables](#5-tables).\n\n## 5. Tables\n")
        anchor = doc.paragraphs[1]._p.find(qn("w:hyperlink")).get(qn("w:anchor"))
        names = [b.get(qn("w:name")) for b in doc.element.body.iter(qn("w:bookmarkStart"))]
        self.assertIn(anchor, names)

    def test_ordered_list_keeps_its_start_number(self):
        md = "# T\n\nIntro\n\n5. Five\n6. Six\n"
        doc = convert(md)
        numbering = doc.part.numbering_part.element
        starts = {el.get(qn("w:val")) for el in numbering.iter(qn("w:start"))}
        self.assertIn("5", starts)

    def test_nested_lists_indent_every_level(self):
        doc = convert("# T\n\n- a\n  - b\n    - c\n      - d\n")
        levels = [p._p.pPr.numPr.ilvl.val for p in doc.paragraphs[1:5]]
        self.assertEqual(levels, [0, 1, 2, 3])

    def test_list_item_code_block_and_second_paragraph(self):
        doc = convert("# T\n\n1. Install:\n\n   ```\n   npm i\n   ```\n\n   Then run it.\n")
        self.assertEqual([t for t in all_text(doc) if t][1:4], ["Install:", "npm i", "Then run it."])
        (code,) = boxes(doc)
        indent = int(code._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")))
        self.assertGreater(indent, 0)

    def test_math_becomes_word_equations(self):
        doc = convert("# T\n\nEuler $e^{i\\pi}+1=0$.\n\n$$\n\\frac{a}{b}\n$$\n\nformula(E = mc^2)\n")
        body = doc.element.body
        self.assertEqual(len(list(body.iter(f"{M}oMathPara"))), 2)
        self.assertGreaterEqual(len(list(body.iter(f"{M}oMath"))), 3)
        self.assertEqual(len(list(body.iter(f"{M}f"))), 1)

    def test_prices_are_not_math(self):
        doc = convert("# T\n\nPro costs $8/mo and Team $4/user/mo.\n")
        self.assertEqual(len(list(doc.element.body.iter(f"{M}oMath"))), 0)
        self.assertEqual(doc.paragraphs[1].text, "Pro costs $8/mo and Team $4/user/mo.")

    def test_tables_size_columns_to_their_content(self):
        doc = convert("# T\n\n| # | Description |\n|---|---|\n| 1 | A much longer description of the thing |\n")
        table = doc.tables[0]
        first, second = (col.width for col in table.columns)
        self.assertLess(first * 3, second)
        rows = table._tbl.findall(qn("w:tr"))
        self.assertTrue(all(tr.trPr.find(qn("w:cantSplit")) is not None for tr in rows))
        self.assertIsNotNone(rows[0].trPr.find(qn("w:tblHeader")))

    def test_html_table_is_a_table(self):
        doc = convert("# T\n\n<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>\n")
        self.assertEqual([c.text for c in doc.tables[0].rows[1].cells], ["1", "2"])

    def test_definition_list(self):
        doc = convert("# T\n\nTerm\n: The definition.\n")
        term, definition = doc.paragraphs[1:3]
        self.assertTrue(term.runs[0].bold)
        self.assertGreater(definition.paragraph_format.left_indent, 0)

    def test_private_blocks_and_front_matter_never_render(self):
        doc = convert("---\ntitle: From front matter\n---\n\nIntro.\n\n```private\nSecret plan\n```\n")
        text = all_text(doc)
        self.assertEqual(doc.paragraphs[0].style.name, "Title")
        self.assertEqual(text[0], "From front matter")
        self.assertNotIn("Secret plan", " ".join(text))
        self.assertNotIn("title:", " ".join(text))

    def test_metadata_line_sits_under_the_title(self):
        doc = convert("Author: Jane | Version: 2\n\n# Report\n\nBody.\n")
        self.assertEqual(body_text(doc)[:3], ["Report", "Author: Jane   ·   Version: 2", "Body."])

    def test_page_break_starts_the_next_block_on_a_new_page(self):
        doc = convert('# T\n\nA\n\n<div style="page-break-after: always"></div>\n\nB\n\n\\pagebreak\n\nC\n')
        texts = {p.text: p for p in doc.paragraphs}
        self.assertTrue(texts["B"].paragraph_format.page_break_before)
        self.assertTrue(texts["C"].paragraph_format.page_break_before)
        self.assertFalse(texts["A"].paragraph_format.page_break_before)

    def test_emoji_shortcodes(self):
        doc = convert("# T\n\nShip it :rocket:\n")
        self.assertEqual(doc.paragraphs[1].text, "Ship it 🚀")

    def test_html_comments_are_dropped(self):
        doc = convert("# T\n\n<!-- internal note -->\n\nVisible\n")
        self.assertNotIn("internal note", " ".join(body_text(doc)))


if __name__ == "__main__":
    unittest.main()
