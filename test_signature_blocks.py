import re
import unittest
from io import BytesIO
from unittest import mock

from docx import Document
from docx.oxml.ns import qn

from app.branding import BrandingConfig
from app.converter import convert_md_to_docx
from app.signature_blocks import (
    SignatureEntry,
    date_tag,
    is_signature_info,
    parse_signature_block,
    signature_tag,
)

BLOCK = """```signature
- role: Employee
  name: John Adams
  title: Head of Engineering
- role: Board member
  name: Jón Jónsson
- role: Witness
```
"""

# A4 portrait: a 16 cm text column, so two boxes of 211 pt; the date line
# is 120 pt. See converter.SIGNATURE_AREA for the units.
TAGS = [
    "{{Employee signature;type=signature;role=Employee;width=211;height=38}}",
    "{{Board member signature;type=signature;role=Board member;width=211;height=38}}",
    "{{Employee date;type=datenow;role=Employee;width=120;height=20}}",
    "{{Board member date;type=datenow;role=Board member;width=120;height=20}}",
    "{{Witness signature;type=signature;role=Witness;width=211;height=38}}",
    "{{Witness date;type=datenow;role=Witness;width=120;height=20}}",
]


def convert(md: str, **kwargs) -> Document:
    return Document(BytesIO(convert_md_to_docx(md, **kwargs)))


def texts(element) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in element.iter(qn("w:p"))]


def tags(doc: Document) -> list[str]:
    return [t for t in texts(doc.element.body) if "{{" in t]


def signature_table(doc: Document):
    """The two-across table of boxes, inside its borderless outer box."""
    (outer,) = doc.tables
    (grid,) = outer.cell(0, 0).tables
    return outer, grid


class ParseSignatureBlockTests(unittest.TestCase):
    def test_spec_example(self):
        self.assertEqual(parse_signature_block(BLOCK.split("\n", 1)[1].rsplit("```", 1)[0]), [
            SignatureEntry("Employee", "John Adams", "Head of Engineering"),
            SignatureEntry("Board member", "Jón Jónsson"),
            SignatureEntry("Witness"),
        ])

    def test_values_are_trimmed_and_lose_one_level_of_quotes(self):
        entries = parse_signature_block(
            "- role:   \"Employee\"  \n  name: ' John '\n  title: \"'Quoted'\"\n"
            "- role: \"Mismatched'\n")
        self.assertEqual(entries, [
            SignatureEntry("Employee", "John", "'Quoted'"),
            SignatureEntry("\"Mismatched'"),
        ])

    def test_keys_are_case_insensitive_and_unknown_keys_ignored(self):
        self.assertEqual(
            parse_signature_block("- ROLE: Employee\n  Email: a@b.is\n  Name: John\n"),
            [SignatureEntry("Employee", "John")])

    def test_any_key_can_start_an_entry_and_indentation_and_blank_lines_are_ignored(self):
        self.assertEqual(
            parse_signature_block("\n- name: John\n\nrole: Employee\n      - role: Witness\n"),
            [SignatureEntry("Employee", "John"), SignatureEntry("Witness")])

    def test_empty_name_and_title_are_dropped(self):
        self.assertEqual(parse_signature_block("- role: A\n  name:\n  title: ''\n"),
                         [SignatureEntry("A")])

    def test_invalid_blocks(self):
        for body in (
            "",
            "  \n",
            "role: Employee\n",                        # no "- " entry
            "Sign here\n- role: Employee\n",           # text before the first entry
            "- name: John\n",                          # no role
            "- role: Employee\n- name: John\n",        # one entry without a role
            "- role:\n",
            "- role: ''\n",
            "- role: " + "x" * 61 + "\n",
            "- role: a;b\n", "- role: a{b\n", "- role: a}b\n", "- role: a=b\n",
            "- role: Employee\n- role: employee\n",   # duplicate, ignoring case
        ):
            with self.subTest(body=body):
                self.assertIsNone(parse_signature_block(body))

    def test_role_length_is_measured_like_javascript(self):
        self.assertIsNotNone(parse_signature_block("- role: " + "x" * 60))
        # 30 emoji are 60 UTF-16 code units (the web's String.length); 31 are 62.
        self.assertIsNotNone(parse_signature_block("- role: " + "😀" * 30))
        self.assertIsNone(parse_signature_block("- role: " + "😀" * 31))

    def test_info_string(self):
        self.assertTrue(is_signature_info("signature"))
        self.assertTrue(is_signature_info(" Signature "))
        self.assertFalse(is_signature_info("signature extra"))
        self.assertFalse(is_signature_info("signatures"))

    def test_tag_text(self):
        self.assertEqual(signature_tag("Board member", width=211, height=38),
                         "{{Board member signature;type=signature;role=Board member;width=211;height=38}}")
        self.assertEqual(date_tag("Board member", width=120, height=20),
                         "{{Board member date;type=datenow;role=Board member;width=120;height=20}}")


class SignatureBlockRenderingTests(unittest.TestCase):
    def test_two_boxes_per_row(self):
        doc = convert("# T\n\n" + BLOCK)
        _, grid = signature_table(doc)
        # Each pair of boxes takes two rows: signature areas with names,
        # then the date lines, so the dates line up.
        self.assertEqual((len(grid.rows), len(grid.columns)), (4, 2))
        cell = lambda r, c: [t for t in texts(grid.cell(r, c)._tc) if t]  # noqa: E731
        self.assertEqual(cell(0, 0), ["John Adams", "Head of Engineering", "Employee"])
        self.assertEqual(cell(0, 1), ["Jón Jónsson", "Board member"])
        self.assertEqual(cell(1, 0), ["Date"])
        self.assertEqual(cell(1, 1), ["Date"])
        self.assertEqual(cell(2, 0), ["Witness"])  # no name: the role is the bold line
        self.assertEqual(cell(2, 1), [])
        self.assertEqual(cell(3, 0), ["Date"])
        borders = grid._tbl.tblPr.find(qn("w:tblBorders"))
        self.assertTrue(all(b.get(qn("w:val")) == "nil" for b in borders))

    def test_box_styles(self):
        brand = BrandingConfig(border_color="#123456", metadata_color="#654321")
        doc = convert("# T\n\n" + BLOCK, brand=brand)
        _, grid = signature_table(doc)
        box = grid.cell(0, 0)
        area, line, name, title, role = box.paragraphs
        bottom = line._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom"))
        self.assertEqual(bottom.get(qn("w:color")), "123456")
        self.assertAlmostEqual(area.paragraph_format.line_spacing.pt + line.paragraph_format.line_spacing.pt, 40)
        self.assertTrue(name.runs[0].bold)
        self.assertEqual(str(title.runs[0].font.color.rgb), "654321")
        self.assertEqual(str(role.runs[0].font.color.rgb), "654321")
        self.assertLess(role.runs[0].font.size, title.runs[0].font.size)
        # The date line is the bottom border of a cell after the "Date" label.
        (date_line,) = grid.cell(1, 0).tables
        blank = date_line.cell(0, 1)._tc
        self.assertEqual(blank.tcPr.find(qn("w:tcBorders")).find(qn("w:bottom")).get(qn("w:color")), "123456")

    def test_block_and_heading_stay_on_one_page(self):
        doc = convert("# T\n\nText.\n\n## Signatures\n\n" + BLOCK)
        outer, grid = signature_table(doc)
        self.assertIsNotNone(outer.rows[0]._tr.trPr.find(qn("w:cantSplit")))
        self.assertTrue(all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in grid.rows))
        heading = outer.cell(0, 0).paragraphs[0]
        self.assertEqual((heading.text, heading.style.name), ("Signatures", "Heading 2"))
        self.assertNotIn("Signatures", [p.text for p in doc.paragraphs])

    def test_heading_page_break_goes_above_the_box(self):
        doc = convert("# T\n\n## Terms\n\nText.\n\n## Signatures\n\n" + BLOCK, page_break_before=2)
        outer, _ = signature_table(doc)
        before = outer._tbl.getprevious()
        self.assertTrue(before.pPr.pageBreakBefore is not None)
        self.assertFalse(outer.cell(0, 0).paragraphs[0].paragraph_format.page_break_before)

    def test_invalid_block_falls_back_to_a_code_block(self):
        doc = convert("# T\n\n```signature\n- name: John\n```\n", signature_fields=True)
        (box,) = doc.tables
        self.assertEqual(len(box.rows), 1)
        self.assertEqual([t for t in texts(box._tbl) if t], ["signature", "- name: John"])
        self.assertEqual(tags(doc), [])

    def test_no_tags_by_default(self):
        plain = convert("# T\n\n" + BLOCK)
        self.assertEqual(tags(plain), [])
        # The tags' lines are there, empty, so the boxes look the same.
        tagged = convert("# T\n\n" + BLOCK, signature_fields=True)
        count = lambda doc: len(list(doc.element.body.iter(qn("w:p"))))  # noqa: E731
        self.assertEqual(count(plain), count(tagged))

    def test_tags_with_signature_fields(self):
        doc = convert("# T\n\n" + BLOCK, signature_fields=True)
        self.assertEqual(tags(doc), TAGS)
        runs = [r for r in doc.element.body.iter(qn("w:r"))
                if "{{" in "".join(t.text or "" for t in r.iter(qn("w:t")))]
        for r in runs:
            rpr = r.rPr
            self.assertEqual(rpr.find(qn("w:color")).get(qn("w:val")), "FFFFFF")
            self.assertEqual(rpr.find(qn("w:sz")).get(qn("w:val")), "2")  # 1 pt
            self.assertIsNotNone(rpr.find(qn("w:noProof")))

    def test_tag_sizes_follow_the_box(self):
        doc = convert("# T\n\n" + BLOCK, signature_fields=True, landscape=True)
        _, grid = signature_table(doc)
        tc = grid.cell(0, 0)._tc
        cell_width = int(tc.tcPr.find(qn("w:tcW")).get(qn("w:w")))
        margin = int(tc.tcPr.find(qn("w:tcMar")).find(qn("w:right")).get(qn("w:w")))
        width = int(re.search(r";width=(\d+);", tags(doc)[0]).group(1))
        self.assertAlmostEqual(width, (cell_width - margin) / 20, delta=1)
        self.assertGreater(width, 300)

    def test_long_role_tags_are_one_paragraph_each(self):
        role = "W" * 60
        doc = convert(f"# T\n\n```signature\n- role: {role}\n```\n", signature_fields=True)
        self.assertEqual([t.split(";")[0] for t in tags(doc)], [f"{{{{{role} signature", f"{{{{{role} date"])


class SignatureFieldsRequestTests(unittest.TestCase):
    """``signature_fields`` reaches the builder from the API."""

    @classmethod
    def setUpClass(cls):
        try:
            from fastapi.testclient import TestClient
        except ImportError as exc:  # httpx missing
            raise unittest.SkipTest(str(exc))
        from app import main

        cls.main = main
        # Without entering the client the lifespan doesn't run: the Word
        # document is built in a thread instead of the process pool.
        cls.client = TestClient(main.app)

    def post(self, path: str, **body):
        response = self.client.post(path, json={"markdown": "# T\n\n" + BLOCK, **body})
        self.assertEqual(response.status_code, 200, response.text)
        return Document(BytesIO(response.content))

    def test_flag_defaults_to_false(self):
        self.assertFalse(self.main.ConvertRequest(markdown="x").signature_fields)
        self.assertEqual(tags(self.post("/convert/docx")), [])

    def test_docx_endpoint(self):
        self.assertEqual(tags(self.post("/convert/docx", signature_fields=True)), TAGS)
        self.assertEqual(tags(self.post("/convert/docx", signature_fields=False)), [])

    def test_pdf_endpoint(self):
        class EchoLibreOffice:
            async def convert(self, docx_bytes):
                return docx_bytes  # the Word document LibreOffice would print

        with mock.patch.object(self.main, "_libreoffice", EchoLibreOffice()):
            self.assertEqual(tags(self.post("/convert/pdf", signature_fields=True)), TAGS)

    def test_convert_md_to_pdf(self):
        from app import pdf_renderer

        with mock.patch.object(pdf_renderer, "convert_oneshot", lambda docx_bytes: docx_bytes):
            docx_bytes = pdf_renderer.convert_md_to_pdf("# T\n\n" + BLOCK, signature_fields=True)
        self.assertEqual(tags(Document(BytesIO(docx_bytes))), TAGS)


if __name__ == "__main__":
    unittest.main()
